#!/usr/bin/env python3
"""Build the R2M manuscript .docx from manuscript_draft.md + figures + tables.

Standard academic manuscript format (Times New Roman 12, double-spaced body,
numbered sections already in the markdown). Figures embedded at the end with
legends; key tables embedded as Word tables.
"""
import re
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Inches, Pt
from docx.oxml.ns import qn

MD = Path('/workspace/r2m/manuscript_draft.md')
OUT = Path('/workspace/r2m/R2M_manuscript_draft_v3.docx')
FIGDIR = Path('/mnt/results/r2m/figures')
TABDIR = Path('/mnt/results/r2m/tables')

FONT = 'Times New Roman'


def set_font(run, size=12, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    r = run._element.rPr
    rfonts = r.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = r.makeelement(qn('w:rFonts'), {})
        r.append(rfonts)
    rfonts.set(qn('w:eastAsia'), FONT)


def add_rich_text(par, text, size=12):
    """Add text with **bold** and *italic* markdown inline parsing."""
    tokens = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)', text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**'):
            run = par.add_run(tok[2:-2]); set_font(run, size, bold=True)
        elif tok.startswith('*') and tok.endswith('*') and len(tok) > 2:
            run = par.add_run(tok[1:-1]); set_font(run, size, italic=True)
        else:
            run = par.add_run(tok); set_font(run, size)


def add_par(doc, text, size=12, bold=False, align=None, space_after=6,
            line_spacing=2.0, style=None):
    par = doc.add_paragraph(style=style)
    if align is not None:
        par.alignment = align
    pf = par.paragraph_format
    pf.space_after = Pt(space_after)
    if line_spacing:
        pf.line_spacing = line_spacing
    if bold:
        run = par.add_run(text); set_font(run, size, bold=True)
    else:
        add_rich_text(par, text, size)
    return par


def add_table_from_csv(doc, csv_path, caption, max_rows=40, font_size=9):
    df = pd.read_csv(csv_path)
    if len(df) > max_rows:
        df = df.head(max_rows)
    add_par(doc, caption, size=11, bold=True, space_after=4, line_spacing=1.0)
    table = doc.add_table(rows=1, cols=len(df.columns))
    table.style = 'Light Grid Accent 1'
    for j, col in enumerate(df.columns):
        cell = table.rows[0].cells[j]
        cell.text = ''
        run = cell.paragraphs[0].add_run(str(col))
        set_font(run, font_size, bold=True)
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for j, val in enumerate(row):
            cells[j].text = ''
            if isinstance(val, float):
                txt = f'{val:.4g}'
            else:
                txt = str(val)
            run = cells[j].paragraphs[0].add_run(txt)
            set_font(run, font_size)
    doc.add_paragraph()


def add_figure(doc, img_path, legend, width=6.5):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = par.add_run()
    run.add_picture(str(img_path), width=Inches(width))
    add_par(doc, legend, size=10, space_after=12, line_spacing=1.0)


def main():
    doc = Document()
    # base style
    style = doc.styles['Normal']
    style.font.name = FONT
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

    lines = MD.read_text().split('\n')
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.startswith('# '):
            add_par(doc, line[2:], size=16, bold=True,
                    align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12, line_spacing=1.5)
        elif line.startswith('## '):
            add_par(doc, line[3:], size=14, bold=True, space_after=8, line_spacing=1.5)
        elif line.startswith('### '):
            add_par(doc, line[4:], size=12, bold=True, space_after=6, line_spacing=1.5)
        elif line.startswith('- '):
            add_par(doc, line[2:], size=12, style='List Bullet')
        else:
            # merge continuation lines into one paragraph
            buf = [line]
            while i + 1 < len(lines) and lines[i + 1].strip() and \
                    not lines[i + 1].startswith(('#', '- ')):
                i += 1
                buf.append(lines[i].rstrip())
            add_par(doc, ' '.join(buf), size=12)
        i += 1

    # ---- Tables ----
    doc.add_page_break()
    add_par(doc, 'Tables', size=14, bold=True, space_after=10, line_spacing=1.5)
    tables = [
        ('table1_baseline_comparison.csv',
         'Table 1. Held-out test-set performance of R2M (seed 42) versus baseline methods '
         '(mean per-gene Pearson correlation across 9,835 genes; paired Wilcoxon '
         'signed-rank test versus R2M, Benjamini-Hochberg FDR).'),
        ('table2_multiseed.csv',
         'Table 2. Multi-seed stability of the main R2M model (three independent '
         'training runs, seeds 42-44; mean +/- SD on the held-out test set).'),
        ('table3_per_cancer.csv',
         'Table 3. Per-cancer-type performance on the held-out test set: number of '
         'test samples, per-sample correlation, and within-cancer per-gene correlation.'),
        ('table4_enrichment_top15.csv',
         'Table 4. Functional enrichment (g:Profiler, FDR < 0.05) of the 500 most '
         'predictable genes; top 15 terms by FDR are shown (128 terms in full '
         'supplementary file).'),
        ('table5_clustering_utility.csv',
         'Table 5. Clustering agreement (NMI/ARI, mean +/- SD over 5 K-means runs) '
         'between cancer-type labels and test-set clusters from true versus predicted '
         'methylation profiles.'),
        ('table6_external_validation.csv',
         'Table 6. External validation on two independent meningioma cohorts with '
         'paired RNA-seq and methylation data (GSE189672/GSE189673, n = 110; '
         'GSE183653/GSE183656, n = 185). R2M used the mean cancer-type embedding '
         '(primary) or the GBM embedding (sensitivity); the train-mean null is a '
         'constant predictor given by the TCGA training-set mean methylation profile. '
         'Centered per-sample r subtracts the cohort mean profile per gene before '
         'correlation, isolating sample-specific signal.'),
        ('table_ablation_v2.csv',
         'Table S2. Ablation analysis on the held-out test set. Main model uses '
         'learnable gene identity embeddings; variants replace them with static '
         'Nucleotide Transformer sequence priors, or remove the cancer-type embedding, '
         'the expert heads, or CLIP contrastive pretraining.'),
        ('table_loco_v2.csv',
         'Table S3. Leave-one-cancer-out zero-shot generalization.'),
    ]
    for fname, cap in tables:
        p = TABDIR / fname
        if p.exists():
            add_table_from_csv(doc, p, cap)

    # ---- Figures ----
    doc.add_page_break()
    add_par(doc, 'Figure Legends', size=14, bold=True, space_after=10, line_spacing=1.5)
    figs = [
        ('/mnt/results/fig1_model_schematic.png',
         'Figure 1. R2M architecture. RNA expression (Top-4000 HVGs) and methylation '
         'profiles are encoded by a CLIP-style dual-tower sample encoder (methylation '
         'tower used during contrastive pretraining only); learnable gene identity '
         'embeddings are encoded by a residual gene encoder; a FiLM-conditioned '
         'bilinear global decoder with cancer-type embedding produces genome-wide '
         'methylation predictions, refined by Top-500/Top-100 expert residual heads. '
         'In the compared variant, learnable gene embeddings are replaced by static '
         'Nucleotide Transformer sequence priors.'),
        ('fig2_test_performance_v2.png',
         'Figure 2. Held-out test-set performance of R2M. (A) Distribution of per-gene '
         'Pearson correlations. (B) Mean correlation on Top-100/500/1000 predictable '
         'gene sets. (C) Distribution of per-sample correlations between predicted and '
         'measured methylation profiles, with the constant train-mean null predictor '
         'shown for reference; the centered per-sample correlation (cohort mean '
         'subtracted per gene) isolates sample-specific signal. (D) Most predictable '
         'genes.'),
        ('fig3_baseline_comparison.png',
         'Figure 3. R2M versus baseline methods on the held-out test set.'),
        ('fig4_ablation_v2.png',
         'Figure 4. Ablation analysis on the held-out test set: replacing the '
         'learnable gene identity embedding with the static Nucleotide Transformer '
         'sequence prior, or removing the cancer-type embedding, the expert heads, '
         'or CLIP contrastive pretraining.'),
        ('fig5_biology.png',
         'Figure 5. Biological determinants of gene-level predictability and functional '
         'enrichment of the most predictable genes.'),
        ('fig6_per_cancer_loco_v2.png',
         'Figure 6. (A) Per-cancer per-sample correlation on the held-out test set '
         '(y-axis truncated at 0.85 to show the 0.92-0.98 range). '
         '(B) Leave-one-cancer-out zero-shot generalization for LUAD: in-distribution '
         'versus zero-shot per-sample and per-gene correlation.'),
        ('fig7_clustering_umap.png',
         'Figure 7. Cancer-type structure in expression and methylation space (UMAP of '
         'held-out test samples colored by cancer type). UMAP axes are independent '
         'across panels and not directly comparable.'),
        ('fig8_external_validation.png',
         'Figure 8. External validation on two independent meningioma cohorts. '
         '(A) Predicted versus observed cohort-mean methylation profile across 9,835 '
         'genes. (B) Per-sample Pearson correlation of R2M versus the constant '
         'train-mean null predictor. (C) Centered per-sample correlation '
         '(sample-specific signal) for the in-distribution TCGA test set, the '
         'no-cancer-embedding ablation, LOCO-LUAD zero-shot prediction, and the two '
         'external cohorts.'),
    ]
    for fname, legend in figs:
        p = Path(fname) if fname.startswith('/') else FIGDIR / fname
        if p.exists():
            add_figure(doc, p, legend)
        else:
            add_par(doc, f'[MISSING FIGURE: {p.name}] ' + legend, size=10,
                    space_after=12, line_spacing=1.0)

    doc.save(OUT)
    print('saved', OUT)


if __name__ == '__main__':
    main()
